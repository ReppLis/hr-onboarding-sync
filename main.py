from load import load_data, load_data_new
from updating import updating
from sort_date import sort_date
from style_sheet import style_data_sort, coloring_cell
from openpyxl import load_workbook
import shutil
import os
from dotenv import load_dotenv
from cell_green import cell_green


def update_sheet(set_update_sheets, del_name, green_cell, source='outside'):
    for sheet in set_update_sheets:
        df_res = load_data_new(os.getenv('PATH_SAVE'), sheet, source=source)
        sort_date(df_res, sheet, del_name, source=source)

        wb_res = load_workbook(os.getenv('PATH_SAVE'))
        style_data_sort(wb_res[sheet], source=source)
        coloring_cell(wb_res, sheet, green_cell)


def main():
    load_dotenv()
    shutil.copy(os.getenv('PATH_SAVE'), os.getenv('PATH_COPY_SAVE'))
    shutil.copy(os.getenv('PATH_LOAD'), os.getenv('PATH_COPY'))
    df = load_data(os.getenv('PATH_COPY'))
    wb = load_workbook(os.getenv('PATH_SAVE'))

    green_cell = cell_green(wb)
    set_update_sheets_outside, del_name_outside = updating(df, wb, green_cell)
    set_update_sheets_inside, del_name_inside = updating(df, wb, green_cell, status='inside')
    wb.save(os.getenv('PATH_SAVE'))

    update_sheet(set_update_sheets_outside, del_name_outside, green_cell)
    update_sheet(set_update_sheets_inside, del_name_inside, green_cell, source='inside')
    print('Данные обновлены')


if __name__ == '__main__':
    main()



