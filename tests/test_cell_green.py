from openpyxl import Workbook
import datetime as dt
from cell_green import cell_green
from openpyxl.styles import PatternFill
from const import THIS_MONTH, CALENDAR_MONTH

def test_cell_green():
    wb = Workbook()
    ws = wb.active
    sheet_name=f'{CALENDAR_MONTH[THIS_MONTH.month]} {THIS_MONTH.year}'
    ws.title = sheet_name

    ws["A2"] = 'Serg'
    ws["A3"] = 'Rita'
    ws["A4"] = 'Dima'

    ws["B2"] = dt.datetime.strptime("16.09.2026", "%d.%m.%Y")
    ws["B3"] = dt. datetime.strptime('18.09.2026', '%d.%m.%Y')
    ws["B4"] = dt. datetime.strptime('20.09.2026', '%d.%m.%Y')

    ws["C2"] = dt.datetime.strptime('17.09.2026', '%d.%m.%Y')
    ws["C3"] = dt.datetime.strptime('23.09.2026', '%d.%m.%Y')
    ws["C4"] = dt.datetime.strptime('28.09.2026', '%d.%m.%Y')
    green = PatternFill(fill_type="solid", start_color="00FF00", end_color="00FF00")

    ws["B2"].fill = green
    ws["C2"].fill = green

    result = cell_green(wb)

    assert ws["A2"].value in result[sheet_name]
    assert ws["A4"].value not in result[sheet_name]
    assert dt.datetime.strptime("16.09.2026", "%d.%m.%Y") in result[sheet_name][ws["A2"].value]
