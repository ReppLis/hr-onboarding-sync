from const import BORDER
import os
from openpyxl.styles import Alignment
from openpyxl.styles import PatternFill


def sheet_format(sheet, source = 'outside'):
    sheet.column_dimensions["A"].width = 20
    sheet['A1'].alignment = Alignment(vertical='top')
    if source == 'outside':
        for col in ["B", "C", "D", "E"]:
            sheet.column_dimensions[col].width = 33
            sheet[col+'1'].alignment = Alignment(wrap_text=True, vertical='top')

    elif source == 'inside':
        for col in ["B", "C"]:
            sheet.column_dimensions[col].width = 33
            sheet[col + '1'].alignment = Alignment(wrap_text=True, vertical='top')
    sheet.row_dimensions[1].height = 207


def format_date(date):
    date.number_format = 'DD.MM.YYYY'


def style_data_sort(sheet, source = 'outside'):
    sheet_format(sheet, source)
    for row in range(1, sheet.max_row + 1):
        if source == 'outside':
            for col in ['A', 'B', 'C', 'D', 'E']:
                if sheet[col + str(row)].value is None:
                    break
                format_date(sheet[col + str(row)])
                sheet[col + str(row)].border = BORDER
        elif source == 'inside':
            for col in ['A', 'B', 'C']:
                if sheet[col + str(row)].value is None:
                    break
                format_date(sheet[col + str(row)])
                sheet[col + str(row)].border = BORDER



def coloring_cell(wb, sheet, cell_green):
    for row in range(2, wb[sheet].max_row + 1):
        if wb[sheet].cell(row,1).value in cell_green[sheet].keys():
            name_cell = wb[sheet].cell(row,1)

            for col in range(1, wb[sheet].max_column + 1):
                if wb[sheet].cell(row,col).value in cell_green[sheet][name_cell.value]:
                    wb[sheet].cell(row, col).fill = PatternFill("solid", fgColor="C6E0B4")

    wb.save(os.getenv('PATH_SAVE'))


# def green_cell(wb, sheet, cell_green):
#     for row, col in cell_green[sheet.title]:
#         sheet.cell(row,col).fill = PatternFill("solid", fgColor="C6E0B4")
#     wb.save(os.getenv('PATH_SAVE'))

